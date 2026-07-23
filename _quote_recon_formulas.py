"""Read-only: XLSX-экспорт CRM (формулы!) → формулы ключевых ячеек «Календаря бронирования»."""
import urllib.request

SHEET_ID = "1sL-rw0klRcJKtWKpgIzLVge6U1GswtBUacJ_jo0JgL0"
url = "https://docs.google.com/spreadsheets/d/" + SHEET_ID + "/export?format=xlsx"
req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
with urllib.request.urlopen(req, timeout=120) as r:
    data = r.read()
print("downloaded bytes:", len(data))
with open("/root/turbobaby-manager-bot/_quote_crm_export.xlsx", "wb") as f:
    f.write(data)

try:
    import openpyxl
except ImportError:
    print("NO openpyxl")
    raise SystemExit(0)

wb = openpyxl.load_workbook("/root/turbobaby-manager-bot/_quote_crm_export.xlsx", data_only=False)
print("sheets:", wb.sheetnames)
ws = None
for n in wb.sheetnames:
    if "алендар" in n:
        ws = wb[n]
        break
print("calendar tab:", ws.title if ws else None, "dims:", ws.dimensions if ws else "-")
for addr in ["F2", "F3", "G3", "H3", "I3", "J3", "A4", "D4", "E4", "F4", "J4",
             "A5", "D5", "E5", "F5", "J5", "J6", "J7", "J20"]:
    c = ws[addr]
    print(addr, "=>", repr(c.value)[:800])
