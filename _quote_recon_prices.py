"""Read-only: формулы и значения вкладки «цены альт» из локального xlsx-экспорта."""
import openpyxl

P = "/root/turbobaby-manager-bot/_quote_crm_export.xlsx"
wbf = openpyxl.load_workbook(P, data_only=False)
wbv = openpyxl.load_workbook(P, data_only=True)
wsf = wbf["цены альт"]
wsv = wbv["цены альт"]
print("dims:", wsf.dimensions)
for row in wsf.iter_rows(min_row=1, max_row=min(wsf.max_row, 60)):
    for c in row:
        f = c.value
        v = wsv[c.coordinate].value
        if f is None and v is None:
            continue
        print(c.coordinate, "| F:", repr(f)[:500], "| V:", repr(v)[:120])
