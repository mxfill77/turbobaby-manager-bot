#!/usr/bin/env python3
# READ-ONLY: export MANAGER workbook via Drive API (live), read Z3:AB15 / K3:M15 of Календарь tab.
import json, urllib.parse, urllib.request, io

SPREADSHEET = '1sL-rw0klRcJKtWKpgIzLVge6U1GswtBUacJ_jo0JgL0'
TAB = 'Календарь бронирования'
tok = json.load(open('/root/.clasprc.json'))['tokens']['default']

def refresh():
    data = urllib.parse.urlencode({
        'client_id': tok['client_id'], 'client_secret': tok['client_secret'],
        'refresh_token': tok['refresh_token'], 'grant_type': 'refresh_token',
    }).encode()
    req = urllib.request.Request('https://oauth2.googleapis.com/token', data=data)
    with urllib.request.urlopen(req, timeout=30) as r:
        return json.load(r)['access_token']

at = refresh()
url = ('https://www.googleapis.com/drive/v3/files/%s/export?mimeType='
       'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
       % SPREADSHEET)
req = urllib.request.Request(url, headers={'Authorization': 'Bearer ' + at})
try:
    with urllib.request.urlopen(req, timeout=90) as r:
        blob = r.read()
except urllib.error.HTTPError as e:
    print('HTTPError', e.code, e.read().decode('utf-8','replace')[:400]); raise SystemExit

print('exported bytes:', len(blob))
import openpyxl
wb = openpyxl.load_workbook(io.BytesIO(blob), data_only=True, read_only=True)
print('tabs:', wb.sheetnames)
if TAB not in wb.sheetnames:
    print('!! tab not found'); raise SystemExit
ws = wb[TAB]
def dump(a1a, a1b, rows):
    print('=== %s:%s ===' % (a1a, a1b))
    for row in ws[a1a:a1b]:
        print([c.value for c in row])
dump('H3', 'J3', 1)
dump('K3', 'M15', 13)
dump('Z3', 'AB15', 13)
