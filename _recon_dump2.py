"""Read-only: памятка целиком; цены альт колонки J..Z; Календарь K..T (где уже блок капов)."""
import openpyxl

def colname(i):
    s = ""
    while i:
        i, rem = divmod(i - 1, 26)
        s = chr(65 + rem) + s
    return s

wbv = openpyxl.load_workbook("/root/turbobaby-manager-bot/_recon_crm.xlsx", data_only=True)
wbf = openpyxl.load_workbook("/root/turbobaby-manager-bot/_recon_crm.xlsx", data_only=False)

print("############### ПАМЯТКА (все ячейки) ###############")
ws = wbv["памятка"]
for ri, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row), start=1):
    for c in row:
        if c.value not in (None, ""):
            print("%s%d = %r" % (colname(c.column), ri, c.value))

print("\n############### ЦЕНЫ АЛЬТ: колонки J..Z, строки 1..32 (есть ли что-то справа) ###############")
ws = wbv["цены альт"]
found = False
for ri in range(1, 33):
    for ci in range(10, 27):  # J..Z
        c = ws.cell(row=ri, column=ci)
        if c.value not in (None, ""):
            print("%s%d = %r" % (colname(ci), ri, c.value))
            found = True
if not found:
    print("  J..Z (стр 1-32) ПУСТО")

print("\n############### КАЛЕНДАРЬ: колонки K..T строки 1..16 (текущий блок капов) ###############")
ws = wbv["Календарь бронирования"]
wf = wbf["Календарь бронирования"]
for ri in range(1, 17):
    cells = []
    for ci in range(11, 21):  # K..T
        c = ws.cell(row=ri, column=ci)
        if c.value not in (None, ""):
            cells.append("%s%d=%r" % (colname(ci), ri, c.value))
    if cells:
        print(" | ".join(cells)[:600])
