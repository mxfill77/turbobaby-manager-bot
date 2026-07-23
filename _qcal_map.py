"""Read-only: карта листа «Календарь бронирования» из ЛОКАЛЬНОГО XLSX-экспорта (формулы).
Ничего не качает, ничего не пишет — читает _quote_crm_export.xlsx от 03.07 12:38."""
import openpyxl

wb = openpyxl.load_workbook("/root/turbobaby-manager-bot/_quote_crm_export.xlsx", data_only=False)
ws = None
for n in wb.sheetnames:
    if "алендар" in n:
        ws = wb[n]
        break
print("TAB:", repr(ws.title), "| dims:", ws.dimensions, "| max_row:", ws.max_row, "| max_col:", ws.max_column)

print("\n=== Строки 1-3 (заголовки/вводные), колонки A..N ===")
for row in range(1, 4):
    for col in range(1, 15):
        c = ws.cell(row=row, column=col)
        if c.value is not None:
            print(f"{c.coordinate} => {repr(c.value)[:300]}")

print("\n=== Строки 4-7 (первые строки данных), A..J ===")
for row in range(4, 8):
    for col in range(1, 11):
        c = ws.cell(row=row, column=col)
        if c.value is not None:
            print(f"{c.coordinate} => {repr(c.value)[:600]}")

print("\n=== Ключевые ячейки ===")
for addr in ["F2", "F3", "G3", "H2", "H3", "I2", "I3", "J2", "J3", "F5", "J4", "J5", "J20", "D4", "E4"]:
    print(addr, "=>", repr(ws[addr].value)[:800])

print("\n=== Есть ли разметка занятости по датам (колонки K+ / строки ниже)? ===")
extra = []
for row in range(1, min(ws.max_row, 60) + 1):
    for col in range(11, ws.max_column + 1):
        c = ws.cell(row=row, column=col)
        if c.value is not None:
            extra.append(f"{c.coordinate} => {repr(c.value)[:200]}")
print("непустых ячеек в K+ (строки 1-60):", len(extra))
for e in extra[:25]:
    print(e)

print("\n=== Хвост листа (строки 41-52, A..J) ===")
for row in range(41, min(ws.max_row, 52) + 1):
    vals = []
    for col in range(1, 11):
        c = ws.cell(row=row, column=col)
        if c.value is not None:
            vals.append(f"{c.coordinate}={repr(c.value)[:120]}")
    if vals:
        print(" | ".join(vals))
